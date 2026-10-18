"""
tools/generar_manual.py
=======================
Genera el manual en .docx y en HTML a partir de una sola fuente por idioma:

    docs/manual/manual_es.md  ─┬─►  Monitor_XAI_Manual_ES_v2.docx
    docs/manual/manual_en.md  ─┼─►  Monitor_XAI_Manual_EN_v2.docx
                               └─►  templates/manual_v2.html  (pestañas ES/EN, ruta /manual-v2)

Así la web y los Word no vuelven a desfasarse: se edita el .md y se regenera.
Los manuales anteriores (Monitor_XAI_Manual_ES.docx, templates/manual.html) no
se tocan.

Requiere python-docx (requirements-dev.txt):
    python tools/generar_manual.py

Admite el subconjunto de Markdown que usan los manuales: # ## ###, párrafos,
listas (- y 1.), tablas, bloques ``` y citas (>), con **negrita**, `código` y
enlaces [texto](url).
"""
import html
import os
import re
import sys

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FUENTES = {"es": os.path.join(RAIZ, "docs", "manual", "manual_es.md"),
           "en": os.path.join(RAIZ, "docs", "manual", "manual_en.md")}
DOCX = {"es": os.path.join(RAIZ, "Monitor_XAI_Manual_ES_v2.docx"),
        "en": os.path.join(RAIZ, "Monitor_XAI_Manual_EN_v2.docx")}
HTML_SALIDA = os.path.join(RAIZ, "templates", "manual_v2.html")


# ── Parser ────────────────────────────────────────────────────────────────────
def parsear(texto: str) -> list:
    """Devuelve bloques: (tipo, contenido). Tipos: h1 h2 h3 p ul ol table code quote."""
    bloques, lineas, i = [], texto.splitlines(), 0
    while i < len(lineas):
        l = lineas[i]
        if not l.strip():
            i += 1
            continue
        if l.startswith("```"):
            j = i + 1
            while j < len(lineas) and not lineas[j].startswith("```"):
                j += 1
            bloques.append(("code", "\n".join(lineas[i + 1:j])))
            i = j + 1
            continue
        m = re.match(r"^(#{1,3}) (.*)", l)
        if m:
            bloques.append((f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
            continue
        if l.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                celdas = [c.strip() for c in lineas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in celdas):
                    filas.append(celdas)
                i += 1
            bloques.append(("table", filas))
            continue
        for tipo, patron in (("ul", r"^- (.*)"), ("ol", r"^\d+\. (.*)"), ("quote", r"^> ?(.*)")):
            if re.match(patron, l):
                items = []
                while i < len(lineas) and re.match(patron, lineas[i]):
                    items.append(re.match(patron, lineas[i]).group(1))
                    i += 1
                bloques.append((tipo, items if tipo != "quote" else " ".join(items)))
                break
        else:
            parrafo = []
            while i < len(lineas) and lineas[i].strip() and not re.match(r"^(#|\||- |\d+\. |> |```)", lineas[i]):
                parrafo.append(lineas[i].strip())
                i += 1
            bloques.append(("p", " ".join(parrafo)))
    return bloques


_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def tokens(texto: str) -> list:
    """Divide en (tipo, texto[, url]) con tipo: txt, b, code, link."""
    out = []
    for parte in _INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**"):
            out.append(("b", parte[2:-2]))
        elif parte.startswith("`"):
            out.append(("code", parte[1:-1]))
        elif parte.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", parte)
            out.append(("link", m.group(1), m.group(2)))
        else:
            out.append(("txt", parte))
    return out


# ── HTML ──────────────────────────────────────────────────────────────────────
def inline_html(texto: str) -> str:
    out = []
    for t in tokens(texto):
        if t[0] == "b":
            out.append(f"<strong>{html.escape(t[1])}</strong>")
        elif t[0] == "code":
            out.append(f'<code class="inline">{html.escape(t[1])}</code>')
        elif t[0] == "link":
            out.append(f'<a href="{html.escape(t[2])}" target="_blank">{html.escape(t[1])}</a>')
        else:
            esc = html.escape(t[1])
            out.append(re.sub(r"(https?://[^\s<]+)", r'<a href="\1" target="_blank">\1</a>', esc))
    return "".join(out)


def bloques_html(bloques: list, lang: str) -> tuple:
    """Devuelve (título, toc_html, cuerpo_html)."""
    titulo, toc, cuerpo, n, abierta = "", [], [], 0, False
    for tipo, c in bloques:
        if tipo == "h1":
            titulo = c
            continue
        if tipo == "h2":
            if abierta:
                cuerpo.append("</section>")
            n += 1
            ancla = f"{lang}-{n}"
            corto = re.sub(r"^\d+\.\s*", "", c).split(" — ")[0].split(":")[0]
            toc.append(f'<a href="#{ancla}">{n}. {html.escape(corto)}</a>')
            cuerpo.append(f'<section class="doc-section" id="{ancla}"><h2 class="manual-h2">{inline_html(c)}</h2>')
            abierta = True
        elif not abierta:
            continue  # lo previo a la primera sección va en la portada (generar_html)
        elif tipo == "h3":
            cuerpo.append(f'<h3 class="sec-h3">{inline_html(c)}</h3>')
        elif tipo == "p":
            cuerpo.append(f'<p class="manual-p">{inline_html(c)}</p>')
        elif tipo in ("ul", "ol"):
            items = "".join(f"<li>{inline_html(x)}</li>" for x in c)
            cuerpo.append(f'<{tipo} class="manual-list">{items}</{tipo}>')
        elif tipo == "quote":
            # Citas: la de la sección 1 es informativa, el aviso legal no lleva negrita inicial, el resto advierte
            clase = "legal" if not c.startswith("**") else ("info" if n == 1 else "warn")
            cuerpo.append(f'<div class="callout {clase}"><p>{inline_html(c)}</p></div>')
        elif tipo == "code":
            cuerpo.append(f'<pre class="codeblock">{html.escape(c)}</pre>')
        elif tipo == "table":
            cab = "".join(f"<th>{inline_html(x)}</th>" for x in c[0])
            filas = "".join("<tr>" + "".join(f"<td>{inline_html(x)}</td>" for x in f) + "</tr>" for f in c[1:])
            cuerpo.append(f'<div class="tabla-wrap"><table class="manual-table"><thead><tr>{cab}</tr></thead>'
                          f"<tbody>{filas}</tbody></table></div>")
    if abierta:
        cuerpo.append("</section>")
    return titulo, "".join(toc), "\n".join(cuerpo)


PLANTILLA = """<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Manual de Usuario / User Manual - Monitor XAI</title>
    <link rel="stylesheet" href="/static/style.css">
    <!-- Generado por tools/generar_manual.py a partir de docs/manual/*.md: no editar a mano. -->
    <style>
        .lang-tabs {{ display:flex; border-bottom:1px solid var(--border); margin-bottom:20px; }}
        .lang-tab-btn {{ background:none; border:none; padding:12px 22px; font-size:.9rem; font-weight:600;
            color:var(--text-muted); cursor:pointer; border-bottom:2px solid transparent; font-family:var(--sans); }}
        .lang-tab-btn:hover {{ color:var(--text-primary); }}
        .lang-tab-btn.active {{ color:var(--accent-cyan); border-bottom-color:var(--accent-cyan); }}
        .lang-content {{ display:none; }}
        .lang-content.active {{ display:block; }}
        .manual-toolbar {{ display:flex; justify-content:space-between; align-items:center; flex-wrap:wrap; gap:12px; margin-bottom:18px; }}
        .manual-toc {{ display:flex; flex-wrap:wrap; gap:8px; margin-bottom:24px; }}
        .manual-toc a {{ font-size:.78rem; color:var(--text-secondary); text-decoration:none; background:rgba(255,255,255,0.03);
            border:1px solid var(--border); border-radius:999px; padding:6px 14px; }}
        .manual-toc a:hover {{ color:var(--accent-cyan); border-color:var(--border-bright); }}
        .manual-portada {{ color:var(--text-secondary); margin-bottom:18px; line-height:1.7; }}
        .manual-h2 {{ font-size:1.3rem; color:var(--text-primary); margin:8px 0 10px; scroll-margin-top:20px; }}
        .doc-section h3.sec-h3 {{ font-size:1rem; color:var(--text-primary); margin:18px 0 8px; }}
        .manual-p {{ color:var(--text-secondary); line-height:1.7; margin-bottom:12px; }}
        .manual-list {{ padding-left:22px; color:var(--text-secondary); line-height:1.8; margin-bottom:12px; }}
        .manual-list strong, .manual-p strong {{ color:var(--text-primary); }}
        .tabla-wrap {{ overflow-x:auto; margin:10px 0 16px; }}
        .manual-table {{ width:100%; border-collapse:collapse; font-size:.86rem; }}
        .manual-table th {{ text-align:left; padding:8px; color:var(--text-primary); border-bottom:1px solid var(--border-bright); }}
        .manual-table td {{ padding:8px; color:var(--text-secondary); border-bottom:1px solid var(--border); vertical-align:top; }}
        .callout {{ border-radius:10px; padding:16px 18px; margin:16px 0; font-size:.88rem; line-height:1.7; border:1px solid var(--border); }}
        .callout.info {{ background:rgba(6,182,212,0.06); border-left:3px solid var(--accent-cyan); color:var(--text-secondary); }}
        .callout.warn {{ background:rgba(245,158,11,0.07); border-left:3px solid var(--accent-amber); color:var(--text-secondary); }}
        .callout.legal {{ background:rgba(239,68,68,0.06); border-left:3px solid var(--accent-red); color:var(--text-secondary); }}
        .callout strong {{ color:var(--text-primary); }}
        code.inline {{ font-family:var(--mono); background:rgba(255,255,255,0.05); padding:2px 6px; border-radius:4px; font-size:.85em; color:var(--accent-cyan); }}
        pre.codeblock {{ background:var(--bg-secondary); border:1px solid var(--border); border-radius:8px; padding:14px 16px;
            overflow-x:auto; font-family:var(--mono); font-size:.8rem; color:var(--text-primary); margin:12px 0; line-height:1.6; }}
        .doc-section a {{ color:var(--accent-cyan); }}
    </style>
</head>
<body>
    <header>
        <div class="header-inner">
            <div class="header-brand">
                <h1>⚖️ Monitor XAI <span class="badge">v1.0</span></h1>
                <p>Ph.D. Vicente Humberto Monteverde · Algoritmos contra la Corrupción</p>
            </div>
            <nav>
                <a href="https://mapatransparencia-production.up.railway.app" target="_blank">🗺️ Mapa</a>
                <a href="/">📊 Dashboard</a>
                <a href="/licitaciones">🏛️ Licitaciones</a>
                <a href="/analisis-vivo">🚀 Análisis en Vivo</a>
                <a href="/documentacion">📖 Documentación</a>
                <a href="/manual" class="active">📘 Manual</a>
                <a href="/estado">🩺 Estado</a>
                <a href="/docs">🔧 API</a>
            </nav>
        </div>
    </header>
    <main>
        <section class="doc-section" style="padding-bottom:0;">
            <div class="manual-toolbar">
                <h2 style="border:none; margin:0; padding:0;">📘 Manual de Usuario / User Manual</h2>
                <div style="display:flex; gap:10px;">
                    <a href="/api/descargar-manual-v2-es" class="btn-descarga" style="padding:8px 16px; font-size:.78rem;">📥 .docx ES</a>
                    <a href="/api/descargar-manual-v2-en" class="btn-descarga" style="padding:8px 16px; font-size:.78rem;">📥 .docx EN</a>
                </div>
            </div>
            <div class="lang-tabs">
                <button class="lang-tab-btn active" data-lang="es" onclick="switchLang('es')">🇦🇷 Español</button>
                <button class="lang-tab-btn" data-lang="en" onclick="switchLang('en')">🇬🇧 English</button>
            </div>
        </section>
{{% raw %}}
{secciones}
{{% endraw %}}
    </main>
    <script>
        function switchLang(lang) {{
            document.querySelectorAll('.lang-tab-btn').forEach(b => b.classList.toggle('active', b.dataset.lang === lang));
            document.querySelectorAll('.lang-content').forEach(c => c.classList.toggle('active', c.id === 'lang-' + lang));
            document.documentElement.lang = lang;
        }}
        // #en, #en-3, etc. abren la versión en inglés (y van a la sección)
        if (location.hash.startsWith('#en')) {{
            switchLang('en');
            // después de "load": si no, el navegador restaura la posición previa al recargar
            window.addEventListener('load', () => {{
                const destino = document.getElementById(location.hash.slice(1));
                if (destino) destino.scrollIntoView();
            }});
        }}
    </script>
</body>
</html>
"""


def generar_html(por_idioma: dict) -> str:
    partes = []
    for lang, bloques in por_idioma.items():
        portada = []
        for tipo, c in bloques:
            if tipo == "h2":
                break
            if tipo == "p":
                portada.append(f"<p>{inline_html(c)}</p>")
        _, toc, cuerpo = bloques_html(bloques, lang)
        activo = " active" if lang == "es" else ""
        partes.append(f'<div id="lang-{lang}" class="lang-content{activo}" lang="{lang}">\n'
                      f'<div class="manual-portada">{"".join(portada)}</div>\n'
                      f'<div class="manual-toc">{toc}</div>\n{cuerpo}\n</div>')
    return PLANTILLA.format(secciones="\n".join(partes))


# ── DOCX ──────────────────────────────────────────────────────────────────────
def generar_docx(bloques: list, ruta: str):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)

    def runs(parrafo, texto):
        for t in tokens(texto):
            if t[0] == "link":
                r = parrafo.add_run(f"{t[1]} ({t[2]})")
            else:
                r = parrafo.add_run(t[1])
            if t[0] == "b":
                r.bold = True
            if t[0] == "code":
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)

    def sombrear(celda, color):
        tc = celda._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tc.append(shd)

    primero_h2 = False
    for tipo, c in bloques:
        if tipo == "h1":
            p = doc.add_heading(c, level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tipo == "h2":
            primero_h2 = True
            doc.add_heading(c, level=1)
        elif tipo == "h3":
            doc.add_heading(c, level=2)
        elif tipo == "p":
            p = doc.add_paragraph()
            runs(p, c)
            if not primero_h2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tipo in ("ul", "ol"):
            for x in c:
                runs(doc.add_paragraph(style="List Bullet" if tipo == "ul" else "List Number"), x)
        elif tipo == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            runs(p, c)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif tipo == "code":
            for linea in c.splitlines():
                r = doc.add_paragraph().add_run(linea)
                r.font.name = "Consolas"
                r.font.size = Pt(9)
        elif tipo == "table":
            t = doc.add_table(rows=len(c), cols=len(c[0]))
            t.style = "Table Grid"
            for i, fila in enumerate(c):
                for j, valor in enumerate(fila):
                    celda = t.cell(i, j)
                    celda.text = ""
                    runs(celda.paragraphs[0], valor)
                    if i == 0:
                        for r in celda.paragraphs[0].runs:
                            r.bold = True
                        sombrear(celda, "D9E2F3")
            doc.add_paragraph()
    doc.save(ruta)


def main():
    por_idioma = {}
    for lang, ruta in FUENTES.items():
        with open(ruta, encoding="utf-8") as f:
            por_idioma[lang] = parsear(f.read())
        generar_docx(por_idioma[lang], DOCX[lang])
        print(f"✅ {os.path.relpath(DOCX[lang], RAIZ)}")
    with open(HTML_SALIDA, "w", encoding="utf-8", newline="\n") as f:
        f.write(generar_html(por_idioma))
    print(f"✅ {os.path.relpath(HTML_SALIDA, RAIZ)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
